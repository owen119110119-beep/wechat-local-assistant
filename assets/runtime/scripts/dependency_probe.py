"""Project-environment checks. Never reads accounts, messages, or browser profiles."""
import argparse
import contextlib
import importlib
import importlib.metadata
import json
import io
import os
from pathlib import Path
import subprocess
import sys
import tempfile

sys.dont_write_bytecode = True


def inspect(root, smoke=False, with_media=False, scratch=None):
    checks = []

    def check(name, modules, action=None):
        try:
            for module in modules:
                importlib.import_module(module)
            if smoke and action:
                action()
            checks.append({'name': name, 'status': 'verified' if smoke else 'installed'})
        except Exception as error:
            reason = 'CJKFontMissing' if str(error) == 'IMAGE_PDF_CJK_FONT_MISSING' else type(error).__name__
            checks.append({'name': name, 'status': 'missing_or_broken', 'reason': reason})

    def versions():
        files = ['requirements-content.txt', 'requirements-media.txt'] + (['requirements-voice.lock.txt'] if with_media else [])
        for name in files:
            for line in (root/'scripts'/name).read_text(encoding='utf-8').splitlines():
                if '==' in line and not line.lstrip().startswith('#'):
                    package, expected = line.strip().split('==', 1)
                    if importlib.metadata.version(package) != expected:raise RuntimeError('Package version mismatch')

    def pdf():
        from image_report import create_report
        from PIL import Image
        source = scratch/'pdf-source.png';Image.new('RGB', (100, 100), 'white').save(source)
        result = create_report({'title': '安装检查', 'overview': '中文课件整理验证',
                                'images': [{'page': 1, 'path': str(source)}],
                                'sections': [{'heading': '要点', 'body': '仅用于检查本机依赖。', 'sourcePages': [1]}]}, scratch/'pdf')
        if not result.get('pdfPath') and not list((scratch/'pdf').glob('*.pdf')):
            raise RuntimeError('PDF not created')

    def office():
        from docx import Document
        from pptx import Presentation
        from openpyxl import Workbook, load_workbook
        doc = Document();doc.add_paragraph('check');doc.save(scratch/'check.docx')
        if Document(scratch/'check.docx').paragraphs[0].text != 'check':raise RuntimeError()
        slides = Presentation();slides.slides.add_slide(slides.slide_layouts[6]);slides.save(scratch/'check.pptx')
        if len(Presentation(scratch/'check.pptx').slides) != 1:raise RuntimeError()
        book = Workbook();book.active['A1'] = 'check';book.save(scratch/'check.xlsx');book.close()
        read = load_workbook(scratch/'check.xlsx')
        try:
            if read.active['A1'].value != 'check':raise RuntimeError()
        finally:read.close()

    def ocr():
        from PIL import Image, ImageDraw, ImageFont
        from rapidocr_onnxruntime import RapidOCR
        image = Image.new('RGB', (420, 100), 'white')
        ImageDraw.Draw(image).text((15, 15), '123456789', font=ImageFont.load_default(size=45), fill='black')
        image.save(scratch/'ocr.png')
        rows, _ = RapidOCR()(str(scratch/'ocr.png'))
        if not rows or '123456789' not in ''.join(row[1].replace(' ', '') for row in rows):raise RuntimeError()

    def browser():
        from playwright.sync_api import sync_playwright
        with sync_playwright() as p:
            from local_tools import browser_options
            options = browser_options(root)
            context = p.chromium.launch(**options, headless=True, chromium_sandbox=True)
            try:
                page = context.new_page();page.set_content('<p>dependency-check</p>')
                if page.locator('p').inner_text() != 'dependency-check':raise RuntimeError()
            finally:context.close()

    def media():
        import av
        from local_tools import ffmpeg_path
        binary = ffmpeg_path()
        subprocess.run([binary, '-version'], check=True, capture_output=True, timeout=15)
        # Encode and reopen a tiny fixture; importing a decoder alone is insufficient.
        path = scratch/'check.mp4'
        with av.open(str(path), 'w') as container:
            stream = container.add_stream('mpeg4', rate=1);stream.width=32;stream.height=32;stream.pix_fmt='yuv420p'
            frame = av.VideoFrame(32, 32, 'yuv420p')
            for plane in frame.planes:plane.update(bytes(plane.buffer_size))
            for packet in stream.encode(frame):container.mux(packet)
            for packet in stream.encode():container.mux(packet)
        with av.open(str(path)) as container:
            if not list(container.decode(video=0)):raise RuntimeError()

    def speech():
        from faster_whisper import WhisperModel
        # No implicit download or fallback to a cloud model.
        WhisperModel(str(root/'.runtime/models/base'), device='cpu', compute_type='int8', local_files_only=True, cpu_threads=1)

    check('article', ['urllib.request', 'html.parser'])
    check('package_versions', [], versions if smoke else None)
    if not smoke:
        try:versions()
        except Exception as error:checks[-1] = {'name': 'package_versions', 'status': 'missing_or_broken', 'reason': type(error).__name__}
    check('web_extract', ['trafilatura', 'chompjs'])
    check('office', ['docx', 'pptx', 'openpyxl'], office)
    check('pdf', ['reportlab', 'pypdf', 'pypdfium2', 'PIL'], pdf)
    check('ocr', ['rapidocr_onnxruntime', 'onnxruntime', 'PIL'], ocr)
    check('browser', ['playwright.sync_api'], browser)
    check('media', ['av', 'yt_dlp', 'imageio_ffmpeg'], media)
    if with_media:
        if not (root/'.runtime/models/base/model.bin').is_file():
            checks.append({'name': 'speech', 'status': 'missing_or_broken', 'reason': 'ModelMissing'})
        else:check('speech', ['faster_whisper', 'huggingface_hub'], speech)
    else:checks.append({'name': 'speech', 'status': 'not_requested'})
    return checks


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--root', required=True);parser.add_argument('--smoke', action='store_true')
    parser.add_argument('--with-media', action='store_true')
    args = parser.parse_args();root=Path(args.root).resolve()
    os.environ.setdefault('PLAYWRIGHT_BROWSERS_PATH', str(root/'.runtime/browsers'))
    with contextlib.redirect_stdout(io.StringIO()):
        if args.smoke:
            temporary = root/'.runtime/setup-checks';temporary.mkdir(parents=True, exist_ok=True)
            with tempfile.TemporaryDirectory(dir=temporary) as folder:
                checks = inspect(root, True, args.with_media, Path(folder))
        else:checks = inspect(root, False, args.with_media)
    print(json.dumps({'checks': checks}, ensure_ascii=False))


if __name__ == '__main__':main()
